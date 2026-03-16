#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
NUMBER_RE = re.compile(r"^(\d+)\.\s+(.*)$")


def set_cell_text(cell, text: str) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    add_inline_runs(paragraph, text.strip())


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_inline_runs(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def add_code_block(document: Document, lines: list[str]) -> None:
    for index, line in enumerate(lines or [""]):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=column_count)
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for index in range(column_count):
        value = rows[0][index] if index < len(rows[0]) else ""
        set_cell_text(header_cells[index], value)
        set_cell_shading(header_cells[index], "D9EAF7")

    for row_values in rows[1:]:
        cells = table.add_row().cells
        for index in range(column_count):
            value = row_values[index] if index < len(row_values) else ""
            set_cell_text(cells[index], value)


def flush_paragraph(document: Document, lines: list[str]) -> None:
    if not lines:
        return
    paragraph = document.add_paragraph()
    add_inline_runs(paragraph, " ".join(line.strip() for line in lines))
    lines.clear()


def flush_table(document: Document, lines: list[str]) -> None:
    if not lines:
        return

    rows: list[list[str]] = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if all(set(cell) <= {"-", ":", " "} for cell in cells):
            continue
        rows.append(cells)

    if rows:
        add_table(document, rows)
    lines.clear()


def export_markdown(source: Path, destination: Path) -> None:
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.size = Pt(10.5)

    paragraph_buffer: list[str] = []
    table_buffer: list[str] = []
    code_buffer: list[str] = []
    in_code_block = False

    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph(document, paragraph_buffer)
            flush_table(document, table_buffer)
            if in_code_block:
                add_code_block(document, code_buffer)
                code_buffer.clear()
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        if stripped.startswith("|"):
            flush_paragraph(document, paragraph_buffer)
            table_buffer.append(line)
            continue

        flush_table(document, table_buffer)

        if not stripped:
            flush_paragraph(document, paragraph_buffer)
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            flush_paragraph(document, paragraph_buffer)
            level = min(len(heading_match.group(1)), 9)
            paragraph = document.add_heading(level=level)
            add_inline_runs(paragraph, heading_match.group(2).strip())
            continue

        bullet_match = BULLET_RE.match(stripped)
        if bullet_match:
            flush_paragraph(document, paragraph_buffer)
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet_match.group(1).strip())
            continue

        number_match = NUMBER_RE.match(stripped)
        if number_match:
            flush_paragraph(document, paragraph_buffer)
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, number_match.group(2).strip())
            continue

        paragraph_buffer.append(line)

    flush_paragraph(document, paragraph_buffer)
    flush_table(document, table_buffer)
    if code_buffer:
        add_code_block(document, code_buffer)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export English-study delivery Markdown docs to .docx")
    parser.add_argument(
        "--source-dir",
        default="/home/gaga/openclaw/var/workspace/projects/english-study/delivery",
        help="Directory containing Markdown delivery docs",
    )
    parser.add_argument(
        "--output-dir",
        default="/home/gaga/openclaw/var/workspace/projects/english-study/delivery/docx",
        help="Directory for generated .docx files",
    )
    args = parser.parse_args()

    source_dir = Path(args.source_dir)
    output_dir = Path(args.output_dir)

    markdown_files = sorted(path for path in source_dir.glob("*.md") if path.is_file())
    if not markdown_files:
        raise SystemExit(f"No Markdown files found under {source_dir}")

    for markdown_file in markdown_files:
        output_file = output_dir / (markdown_file.stem + ".docx")
        export_markdown(markdown_file, output_file)
        print(f"exported {markdown_file.name} -> {output_file.name}")


if __name__ == "__main__":
    main()
